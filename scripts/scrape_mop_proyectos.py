"""
Construye la hoja 'Base MOP' de Bases de dato/BASE MOP.xlsx desde las fichas
de https://proyectos.mop.gob.cl (sin cruces con BIP ni BIDAT).

  1. Lista todos los proyectos del buscador (Default.asp) → cod_p.
  2. Descarga por proyecto la ficha XML ("Versión en XML", fuente principal)
     y la ficha HTML (proyecto.asp), que se usa de respaldo y para comparar
     el Total General. Ambas quedan en caché en Bases de dato/MOP fichas/.
  3. Consulta el servicio ArcGIS que dibuja el mapa de la ficha para contar
     los proyectos geolocalizados (solo se informan, no van al Excel).
  4. Reescribe la hoja 'Base MOP' y elimina las hojas del cruce antiguo
     ('Resumen BIP', 'BIDAT sin duplicados'). La primera vez guarda un
     respaldo del libro original.

Uso:  python scripts/scrape_mop_proyectos.py [--refresh]
      --refresh  vuelve a descargar las fichas aunque estén en caché.
"""
import os
import re
import sys
import time
import shutil
import unicodedata
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, 'Bases de dato')
EXCEL_PATH = os.path.join(DATA_DIR, 'BASE MOP.xlsx')
BACKUP_PATH = os.path.join(DATA_DIR, 'BASE MOP (respaldo cruce BIP).xlsx')
CACHE_DIR = os.path.join(DATA_DIR, 'MOP fichas')

SITE = 'https://proyectos.mop.gob.cl'
# Los enlaces de paginación del sitio (con sqlQuery=...) los bloquea su firewall;
# sin ese parámetro y con un pagesize grande el buscador devuelve todo en una página.
LIST_URL = SITE + '/Default.asp?whichpage=1&buscar=true&pagesize=5000'
FICHA_URL = SITE + '/proyecto.asp?cod_p={cod}'
XML_URL = SITE + '/descargar.asp?file=fichas_xml%5CCodigo{cod}.xml'
ARCGIS_URL = 'https://rest-sit.mop.gob.cl/arcgis/rest/services/INTEROP/EXPLORATORIO_PROYECTOS_VISION/MapServer/{layer}/query'
# Mismo filtro que usa el mapa de la ficha (ide.mop.gob.cl/Framework/data/config.js)
ARCGIS_FILTER = "(ESTADO_EDIT LIKE '2' OR ESTADO_EDIT LIKE '3') AND TIPO LIKE '1'"

HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) CATLEC-Observatorio'}
WORKERS = 4
RETRIES = 4
REMOVE_SHEETS = ['Base MOP', 'Resumen BIP', 'BIDAT sin duplicados']

# Columnas de montos de la tabla "Montos de Inversión Global" (XML → orden de la ficha)
INV_KEYS = ['EJECUTADO_AGNO_ANT', 'AGNO_ACT', 'AGNO_01', 'AGNO_02', 'AGNO_03',
            'AGNO_04', 'AGNO_05', 'Monto_Saldo']

COLUMNS = [
    'CODIGO BIP', 'CODIGO PROYECTO MOP', 'NOMBRE', 'FINANCIAMIENTO', 'PROGRAMA',
    'SERVICIO', 'ETAPA VIGENTE DE PROYECTO', 'REGION', 'PROVINCIA', 'COMUNA',
    'PLAN (ES)', 'DESCRIPCION', 'JUSTIFICACION', 'TOTAL GENERAL INVERSION GLOBAL (M$)',
    'AÑO ESTIMADO DE OPERACION', 'OBS. AÑO OPERACION', 'URL FICHA', 'FUENTE',
]
SITE_SOURCE = 'proyectos.mop.gob.cl'
MANUAL_SOURCE = 'Manual'

# Nombres oficiales de región (el sitio a veces escribe variantes, ej. "DE LOS LAGOS")
REGIONS = [
    'Arica y Parinacota', 'Tarapacá', 'Antofagasta', 'Atacama', 'Coquimbo', 'Valparaíso',
    'Metropolitana de Santiago', "Libertador General Bernardo O'Higgins", 'Maule', 'Ñuble',
    'Biobío', 'La Araucanía', 'Los Ríos', 'Los Lagos',
    'Aysén del General Carlos Ibáñez del Campo', 'Magallanes y de la Antártica Chilena',
]
PLACEHOLDERS = {'', '-', 'no existe informacion'}

session = requests.Session()
session.headers.update(HEADERS)


def _norm(s):
    s = unicodedata.normalize('NFD', str(s)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', s).strip()


def _clean(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def mop_id(val):
    """ID de proyecto como texto: 38417, 38417.0 y '38417' → '38417'; 'M-001' se mantiene."""
    if val is None or (isinstance(val, float) and val != val):
        return ''
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    s = str(val).strip()
    return s[:-2] if re.fullmatch(r'\d+\.0', s) else s


def _region(name):
    n = re.sub(r'^(region )?de ', '', _norm(name))
    for ref in REGIONS:
        if _norm(ref) == n:
            return ref
    return _clean(name)


def _num(s):
    """'53.188.720' (HTML) o '53188720' (XML) → int."""
    s = str(s or '').strip()
    if not s:
        return 0
    if re.fullmatch(r'-?\d{1,3}(\.\d{3})+', s):
        s = s.replace('.', '')
    try:
        return int(round(float(s.replace(',', '.'))))
    except ValueError:
        return 0


def fetch(url, cache_path=None, refresh=False):
    """GET con reintentos; detecta la página de bloqueo del firewall."""
    if cache_path and not refresh and os.path.exists(cache_path) and os.path.getsize(cache_path) > 0:
        with open(cache_path, 'rb') as f:
            return f.read()
    last_err = None
    for attempt in range(RETRIES):
        try:
            r = session.get(url, timeout=60)
            if r.status_code == 404:
                return None
            r.raise_for_status()
            if b'Unauthorized Request Blocked' in r.content[:3000]:
                raise RuntimeError('bloqueado por el firewall del sitio')
            if cache_path:
                with open(cache_path, 'wb') as f:
                    f.write(r.content)
            return r.content
        except Exception as e:
            last_err = e
            time.sleep(2 * (attempt + 1))
    raise RuntimeError(f'{url}: {last_err}')


# ── 1. Listado de proyectos ──────────────────────────────────────────────────
def list_projects(refresh):
    raw = fetch(LIST_URL, os.path.join(CACHE_DIR, '_listado.html'), refresh)
    cods = []
    for c in re.findall(rb'proyecto\.asp\?cod_p=(\d+)', raw):
        c = int(c)
        if c not in cods:
            cods.append(c)
    return cods


# ── 2a. Ficha XML ────────────────────────────────────────────────────────────
def txt_of(el, tag):
    return _clean(el.findtext(tag))


def parse_xml(raw):
    root = ET.fromstring(raw)
    p = root.find('PROYECTO')
    # Fichas rotas del sitio: XML de plantilla sin código ni nombre
    if p is None or not _clean(p.findtext('CODIGO')) or not _clean(p.findtext('NOMBRE')):
        return None

    def txt(path):
        return _clean(p.findtext(path))

    def many(path):
        return [_clean(e.text) for e in p.findall(path) if _norm(e.text or '') not in PLACEHOLDERS]

    etapas = []
    for e in p.findall('INVERSION/ETAPAS'):
        etapas.append((_clean(e.findtext('ETAPA_MIDEPLAN')), [_num(e.findtext(k)) for k in INV_KEYS]))

    indicadores = [(_clean(i.findtext('NOMBRE')), _num(i.findtext('VALOR'))) for i in p.findall('INDICADORES/INDICADOR')]
    return {
        'agno': _num(root.findtext('AGNO_VIGENTE')),
        'bip': txt('CODIGO_BIP'),
        'cod_p': _num(txt('CODIGO')),
        'nombre': txt('NOMBRE'),
        'financiamiento': txt('FINANCIAMIENTO'),
        'programa': txt('PROGRAMA'),
        'servicio': txt('SERVICIO'),
        'etapa': txt('ETAPA_VIGENTE'),
        'region': _region(txt('REGION')),
        'provincias': many('PROVINCIAS/PROVINCIA'),
        'comunas': many('COMUNAS/COMUNA'),
        'planes': many('PLANES/PLANES'),
        'descripcion': txt('DESCRIPCION'),
        'justificacion': txt('JUSTIFICACION'),
        'etapas': etapas,
        'fecha_montos': _clean(p.findtext('INVERSION/ETAPAS/AGNO_MONTO')),
        # Datos adicionales (solo para el resumen)
        'genero': txt('GENERO'),
        'presupuesto': [_num(p.findtext(f'PRESUPUESTO/{k}')) for k in ('LEY', 'DECRETADO', 'EJECUTADO')],
        'rates': [(_num(r.findtext('AGNO')), txt_of(r, 'RATE')) for r in p.findall('RATES/DETALLE_RATE')
                  if _num(r.findtext('AGNO')) > 0],
        'clasificadores': [txt_of(c, 'NOMBRE') for c in p.findall('CLASIFICADORES/CLASIFICADOR')
                           if _norm(txt_of(c, 'NOMBRE')) not in PLACEHOLDERS],
        'indicadores': indicadores,
        'mapa_flag': txt('MAPAS/MAPA').lower() == 'true',
        'imagen': txt('IMAGEN'),
    }


# ── 2b. Ficha HTML (respaldo y control) ──────────────────────────────────────
HTML_LABELS = {
    'codigo bip': 'bip', 'nombre proyecto': 'nombre', 'codigo proyecto': 'cod_p',
    'financiamiento': 'financiamiento', 'programa': 'programa', 'servicio': 'servicio',
    'etapa vigente de proyecto': 'etapa', 'region': 'region', 'provincia(s)': 'provincias',
    'comuna(s)': 'comunas', 'plan(es)': 'planes',
}


def parse_html(raw):
    soup = BeautifulSoup(raw.decode('cp1252', errors='replace'), 'lxml')
    d = {}
    for td in soup.select('td.cabecera_01'):
        label = _norm(td.get_text()).rstrip(':').strip()
        key = HTML_LABELS.get(label)
        nxt = td.find_next_sibling('td')
        if key and nxt is not None and key not in d:
            d[key] = _clean(nxt.get_text())
    if not d.get('cod_p') or not d.get('nombre'):
        return None
    d['cod_p'] = _num(d['cod_p'])
    d['region'] = _region(d.get('region', ''))
    for k in ('provincias', 'comunas', 'planes'):
        d[k] = [s.strip() for s in d.get(k, '').split(' - ') if _norm(s) not in PLACEHOLDERS]

    # Textos: la celda bajo los títulos "Justificación (objetivos)" y "Descripción"
    for key, title in (('justificacion', 'justificacion'), ('descripcion', 'descripcion')):
        d[key] = ''
        for t in soup.select('tr.titulo td'):
            if _norm(t.get_text()).startswith(title):
                nxt_tr = t.find_parent('tr').find_next_sibling('tr')
                if nxt_tr is not None:
                    d[key] = _clean(nxt_tr.get_text())
                break

    # Tabla de montos: año vigente en el título y filas de 9 celdas por etapa
    d['agno'], d['etapas'], d['total_general'] = 0, [], None
    for t in soup.select('tr.titulo td'):
        m = re.search(r'Montos de Inver.*?(\d{4})', t.get_text())
        if m:
            d['agno'] = int(m.group(1))
            table = t.find_parent('table')
            for tr in table.select('tr.datos_reporte'):
                cells = [_clean(c.get_text()) for c in tr.find_all('td', recursive=False)]
                if len(cells) == 9:
                    d['etapas'].append((cells[0], [_num(c) for c in cells[1:]]))
            m_tot = re.search(r'Total General\s*:\s*([\d\.]+)', table.get_text())
            if m_tot:
                d['total_general'] = _num(m_tot.group(1))
            break
    return d


def total_general(etapas):
    return sum(sum(v) for _, v in etapas)


def operation_year(etapas, agno):
    """Primer año sin gasto una vez completado el Total General (fila "Total Año")."""
    tot = [sum(v[i] for _, v in etapas) for i in range(len(INV_KEYS))]
    if not any(tot):
        return None, 'Sin montos informados'
    if tot[-1] > 0:
        return None, f'Saldo pendiente posterior a {agno + 5}'
    last = max(i for i in range(len(tot) - 1) if tot[i] > 0)
    # índice 0 = ejecutado hasta agno-1; índice i = año agno-1+i → operación al año siguiente
    if last == 0:
        return agno, f'Inversión ejecutada antes de {agno}'
    return agno + last, ''


def scrape_one(cod, refresh):
    # Basta con una de las dos versiones de la ficha; solo falla si no hay ninguna
    raw_xml = raw_html = None
    errors = []
    for kind, url in (('xml', XML_URL), ('html', FICHA_URL)):
        try:
            raw = fetch(url.format(cod=cod), os.path.join(CACHE_DIR, f'{cod}.{kind}'), refresh)
        except RuntimeError as e:
            errors.append(str(e))
            continue
        if kind == 'xml':
            raw_xml = raw
        else:
            raw_html = raw
    if not raw_xml and not raw_html:
        raise RuntimeError(' | '.join(errors) or f'cod_p {cod}: ficha no encontrada')
    x = None
    if raw_xml:
        try:
            x = parse_xml(raw_xml)
        except ET.ParseError:
            x = None
    h = parse_html(raw_html) if raw_html else None
    return cod, x, h


# ── 3. Geolocalización (servicio ArcGIS del mapa de la ficha) ────────────────
def geolocated(cods, layer):
    found = set()
    for i in range(0, len(cods), 100):
        chunk = cods[i:i + 100]
        params = {
            'where': f"{ARCGIS_FILTER} AND CODPROYECTO IN ({','.join(map(str, chunk))})",
            'outFields': 'CODPROYECTO', 'returnGeometry': 'false',
            'returnDistinctValues': 'true', 'f': 'json',
        }
        r = session.get(ARCGIS_URL.format(layer=layer), params=params, timeout=60)
        r.raise_for_status()
        data = r.json()
        if 'error' in data:
            raise RuntimeError(data['error'])
        found.update(int(f['attributes']['CODPROYECTO']) for f in data.get('features', []))
    return found


# ── 4. Excel ─────────────────────────────────────────────────────────────────
def manual_rows(site_ids):
    """Filas agregadas a mano en la hoja actual (FUENTE distinta del sitio o vacía)."""
    try:
        wb = load_workbook(EXCEL_PATH, read_only=True)
    except FileNotFoundError:
        return []
    if 'Base MOP' not in wb.sheetnames:
        return []
    it = wb['Base MOP'].iter_rows(values_only=True)
    header = [str(h or '').strip() for h in next(it, [])]
    keep = []
    if 'FUENTE' not in header:
        # Hoja de una versión anterior del script: todas sus filas venían del sitio
        wb.close()
        return keep
    for values in it:
        row = dict(zip(header, values))
        rid = mop_id(row.get('CODIGO PROYECTO MOP'))
        if not rid and not row.get('NOMBRE'):
            continue
        if str(row.get('FUENTE') or '').strip() == SITE_SOURCE:
            continue
        if rid in site_ids:
            print(f"[WARN] Fila manual con ID {rid} repetido en el sitio: se usa la del sitio")
            continue
        out = {c: row.get(c) for c in COLUMNS}
        out['CODIGO PROYECTO MOP'] = int(rid) if rid.isdigit() else rid
        out['FUENTE'] = MANUAL_SOURCE
        keep.append(out)
    wb.close()
    return keep


def write_excel(rows):
    if not os.path.exists(BACKUP_PATH):
        shutil.copy2(EXCEL_PATH, BACKUP_PATH)
        print(f"OK Respaldo del libro original: {BACKUP_PATH}")

    wb = load_workbook(EXCEL_PATH)
    for name in REMOVE_SHEETS:
        if name in wb.sheetnames:
            del wb[name]
    ws = wb.create_sheet('Base MOP', 0)
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r[c] for c in COLUMNS])

    header_fill = PatternFill('solid', fgColor='1B365D')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    widths = {'NOMBRE': 60, 'DESCRIPCION': 70, 'JUSTIFICACION': 70, 'PLAN (ES)': 40,
              'PROGRAMA': 36, 'SERVICIO': 30, 'OBS. AÑO OPERACION': 32, 'URL FICHA': 48}
    for idx, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, 18)
    col_total = COLUMNS.index('TOTAL GENERAL INVERSION GLOBAL (M$)') + 1
    for row in ws.iter_rows(min_row=2, min_col=col_total, max_col=col_total):
        row[0].number_format = '#,##0'
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions
    wb.active = 0
    wb.save(EXCEL_PATH)
    print(f"OK Hoja 'Base MOP' escrita en {EXCEL_PATH} ({len(rows)} proyectos)")


def main():
    refresh = '--refresh' in sys.argv
    os.makedirs(CACHE_DIR, exist_ok=True)

    cods = list_projects(refresh)
    print(f"-> {len(cods)} proyectos en el buscador de {SITE}")

    results, failed = {}, []
    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        futures = [pool.submit(scrape_one, c, refresh) for c in cods]
        for n, fut in enumerate(futures, start=1):
            try:
                cod, x, h = fut.result()
                results[cod] = (x, h)
            except Exception as e:
                failed.append(str(e))
            if n % 100 == 0:
                print(f"   {n}/{len(cods)} fichas")

    rows, extras, diffs, from_html, empty = [], [], [], 0, []
    for cod in cods:
        if cod not in results:
            continue
        x, h = results[cod]
        d = x or h
        if not d:
            empty.append(cod)
            continue
        if not x:
            from_html += 1
        tg = total_general(d['etapas'])
        if h and h.get('total_general') is not None and h['total_general'] != tg:
            diffs.append(f"cod_p {cod}: Total General XML {tg:,} ≠ HTML {h['total_general']:,}")
        if x and h and (_norm(x['etapa']) != _norm(h.get('etapa', '')) or _norm(x['nombre']) != _norm(h.get('nombre', ''))):
            diffs.append(f"cod_p {cod}: nombre/etapa distintos entre XML y HTML")
        year_op, obs = operation_year(d['etapas'], d['agno'])
        rows.append({
            'CODIGO BIP': d['bip'],
            'CODIGO PROYECTO MOP': cod,
            'NOMBRE': d['nombre'],
            'FINANCIAMIENTO': d['financiamiento'],
            'PROGRAMA': d['programa'],
            'SERVICIO': d['servicio'],
            'ETAPA VIGENTE DE PROYECTO': d['etapa'],
            'REGION': d['region'],
            'PROVINCIA': ' - '.join(d['provincias']),
            'COMUNA': ' - '.join(d['comunas']),
            'PLAN (ES)': ' - '.join(d['planes']),
            'DESCRIPCION': d['descripcion'],
            'JUSTIFICACION': d['justificacion'],
            'TOTAL GENERAL INVERSION GLOBAL (M$)': tg,
            'AÑO ESTIMADO DE OPERACION': year_op,
            'OBS. AÑO OPERACION': obs,
            'URL FICHA': FICHA_URL.format(cod=cod),
            'FUENTE': SITE_SOURCE,
            '_agno': d['agno'],
        })
        if x:
            extras.append(x)

    # Fichas desactualizadas en el sitio (XML de otro año vigente): se conservan con aviso
    agno_vigente = max(set(r['_agno'] for r in rows), key=lambda a: sum(r['_agno'] == a for r in rows))
    stale = [r for r in rows if r['_agno'] != agno_vigente]
    for r in stale:
        # Sus montos no son comparables con el año vigente: no se estima el año de operación
        r['AÑO ESTIMADO DE OPERACION'] = None
        r['OBS. AÑO OPERACION'] = f"Ficha del sitio desactualizada (montos de {r['_agno']})"

    # Misma escritura para servicio y programa (algunas fichas antiguas vienen en MAYÚSCULAS)
    for col in ('SERVICIO', 'PROGRAMA'):
        forms = {}
        for r in rows:
            forms.setdefault(_norm(r[col]), []).append(r[col])
        canon = {k: max(set(v), key=lambda f: (v.count(f), f != f.upper())) for k, v in forms.items()}
        for r in rows:
            r[col] = canon[_norm(r[col])]

    manual = manual_rows({mop_id(r['CODIGO PROYECTO MOP']) for r in rows})
    write_excel(rows + manual)
    print(f"Filas manuales conservadas: {len(manual)}")

    # ── Resumen ──────────────────────────────────────────────────────────────
    all_cods = [r['CODIGO PROYECTO MOP'] for r in rows]
    try:
        lines = geolocated(all_cods, 0)
        points = geolocated(all_cods, 1)
        geo = lines | points
        geo_msg = (f"{len(geo)} (solo línea {len(lines - points)}, solo punto {len(points - lines)}, "
                   f"ambos {len(lines & points)}); sin geolocalización: {len(all_cods) - len(geo)}")
    except Exception as e:
        geo_msg = f"[WARN] no se pudo consultar ArcGIS: {e}"

    def count(pred):
        return sum(1 for r in rows if pred(r))

    def count_x(pred):
        return sum(1 for x in extras if pred(x))

    years = {}
    for r in rows:
        key = r['AÑO ESTIMADO DE OPERACION'] or r['OBS. AÑO OPERACION']
        years[key] = years.get(key, 0) + 1

    print('\n=== Resumen fichas MOP ===')
    print(f"Proyectos en la base: {len(rows)} (desde HTML por falta de XML: {from_html})")
    print(f"Fichas con error de descarga: {len(failed)}; fichas vacías: {len(empty)} {empty[:10]}")
    print(f"Con código BIP: {count(lambda r: bool(r['CODIGO BIP']))}")
    print(f"Con descripción: {count(lambda r: bool(r['DESCRIPCION']))}; con justificación: {count(lambda r: bool(r['JUSTIFICACION']))}")
    print(f"Con plan(es): {count(lambda r: bool(r['PLAN (ES)']))}")
    print(f"Total General > 0: {count(lambda r: r['TOTAL GENERAL INVERSION GLOBAL (M$)'] > 0)}")
    print(f"Geolocalizados (mapa ArcGIS de la ficha): {geo_msg}")
    print(f"Marca MAPA=True en el XML: {count_x(lambda x: x['mapa_flag'])}")
    print(f"Con presupuesto MOP (Ley/Decretado/Ejecutado > 0): {count_x(lambda x: any(x['presupuesto']))}")
    print(f"Con RATE (resultado del análisis técnico-económico) informado: {count_x(lambda x: any(r for _, r in x['rates']))}")
    print(f"Con clasificadores (política rural, inclusión, género...): {count_x(lambda x: bool(x['clasificadores']))}")
    for nombre in ('BENEFICIARI', 'COSTO ANUAL EQUIVALENTE', 'TASA INTERNA DE RETORNO', 'VAN SOCIAL'):
        print(f"Con indicador '{nombre}' > 0: "
              f"{count_x(lambda x: any(v > 0 and n.startswith(nombre) for n, v in x['indicadores']))}")
    print(f"Con enfoque de género (GENERO = SI): {count_x(lambda x: _norm(x['genero']) == 'si')}")
    print(f"Fecha de actualización de montos: {sorted({x['fecha_montos'] for x in extras if x['fecha_montos']})[:5]}")
    print(f"Año estimado de operación: {dict(sorted(years.items(), key=lambda kv: str(kv[0])))}")
    print(f"Fichas desactualizadas (año vigente distinto de {agno_vigente}): "
          f"{[r['CODIGO PROYECTO MOP'] for r in stale]}")
    print(f"Diferencias XML/HTML: {len(diffs)}")
    for d in diffs[:20]:
        print('   ' + d)
    for f in failed[:10]:
        print('   [WARN] ' + f)


if __name__ == '__main__':
    main()
